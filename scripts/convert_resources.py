"""Copy a session's resources into the project and convert each one to text.

Usage (from the project root):
    .venv/Scripts/python.exe scripts/convert_resources.py <source_dir> <session_dir> --session 2
    .venv/Scripts/python.exe scripts/convert_resources.py <source_dir> inputs/curso/recursos --shared

For every non-video file in <source_dir>:
    PDF  -> <name>.pdf.md   (PyMuPDF, page by page, with the number of images per page so
                             screenshots are not overlooked)
    XLSX -> <name>.xlsx.md  (one Markdown table per sheet; a cell with a formula shows its
                             cached value followed by the formula)
    MD/TXT -> copied as is
The original is copied next to its text version. A file identical (MD5) to one already in
inputs/curso/recursos/ is not copied again: fuente.md points to the shared copy instead.
In session mode, fuente.md records the video path and duration (ffprobe) and every resource.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import shutil
import subprocess
import sys
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SHARED_DIR = PROJECT_ROOT / "inputs" / "curso" / "recursos"
VIDEO_SUFFIXES = {".mp4", ".mov", ".mkv", ".m4a", ".wav"}
CONVERTED_SUFFIXES = (".pdf.md", ".xlsx.md")


def md5(path: Path) -> str:
    digest = hashlib.md5()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def hms(seconds: float) -> str:
    total = int(seconds)
    return f"{total // 3600}:{total % 3600 // 60:02d}:{total % 60:02d}"


def probe_duration(video: Path) -> float:
    out = subprocess.run(
        ["ffprobe", "-v", "error", "-show_entries", "format=duration",
         "-of", "default=noprint_wrappers=1:nokey=1", str(video)],
        check=True, capture_output=True, text=True,
    ).stdout.strip()
    return float(out)


def pdf_to_markdown(pdf: Path) -> str:
    import pymupdf

    doc = pymupdf.open(pdf)
    parts = [
        f"# {pdf.name}",
        "",
        f"Texto extraído de `{pdf.name}` ({doc.page_count} páginas). Las imágenes no se extraen; "
        "cada página indica cuántas tiene, para no pasarlas por alto.",
    ]
    for number, page in enumerate(doc, start=1):
        images = len(page.get_images(full=True))
        heading = f"## Página {number}" + (f" · {images} imagen(es)" if images else "")
        parts += ["", "---", "", heading, "", page.get_text("text").strip()]
    return "\n".join(parts) + "\n"


def xlsx_to_markdown(path: Path) -> str:
    import openpyxl

    with_formulas = openpyxl.load_workbook(path, data_only=False)
    with_values = openpyxl.load_workbook(path, data_only=True)
    parts = [f"# {path.name}", "", f"Contenido de `{path.name}`. En las celdas con fórmula: valor guardado, luego la fórmula."]
    for sheet in with_formulas.worksheets:
        values = with_values[sheet.title]
        parts += ["", f"## Hoja: {sheet.title} ({sheet.max_row} filas × {sheet.max_column} columnas)", ""]
        letters = [openpyxl.utils.get_column_letter(i) for i in range(1, sheet.max_column + 1)]
        parts.append("| fila | " + " | ".join(letters) + " |")
        parts.append("|---|" + "---|" * len(letters))
        for row in sheet.iter_rows():
            cells = []
            for cell in row:
                raw = cell.value
                if isinstance(raw, str) and raw.startswith("="):
                    cached = values[cell.coordinate].value
                    cells.append(f"{'' if cached is None else cached} `{raw}`")
                else:
                    cells.append("" if raw is None else str(raw).replace("\n", " ").replace("|", "\\|"))
            if any(cells):
                parts.append(f"| {row[0].row} | " + " | ".join(cells) + " |")
    return "\n".join(parts) + "\n"


def convert(original: Path, dest_dir: Path) -> Path | None:
    """Copy the original into dest_dir and write its text version. Returns the text file."""
    dest_dir.mkdir(parents=True, exist_ok=True)
    copied = dest_dir / original.name
    shutil.copy2(original, copied)
    suffix = original.suffix.lower()
    if suffix == ".pdf":
        text_file = dest_dir / f"{original.name}.md"
        text_file.write_text(pdf_to_markdown(original), encoding="utf-8")
        return text_file
    if suffix == ".xlsx":
        text_file = dest_dir / f"{original.name}.md"
        text_file.write_text(xlsx_to_markdown(original), encoding="utf-8")
        return text_file
    return copied  # md/txt: the original is already text


def shared_index() -> dict[str, Path]:
    if not SHARED_DIR.is_dir():
        return {}
    return {md5(p): p for p in SHARED_DIR.iterdir() if p.is_file() and not p.name.endswith(CONVERTED_SUFFIXES)}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source_dir", type=lambda p: Path(p).resolve())
    parser.add_argument("dest_dir", type=lambda p: Path(p).resolve())
    parser.add_argument("--session", type=int, help="session number, for fuente.md")
    parser.add_argument("--shared", action="store_true", help="convert into the shared resources folder (no fuente.md)")
    args = parser.parse_args()
    if not args.shared and args.session is None:
        parser.error("either --session N or --shared is required")

    files = sorted(p for p in args.source_dir.iterdir() if p.is_file())
    videos = [p for p in files if p.suffix.lower() in VIDEO_SUFFIXES]
    resources = [p for p in files if p.suffix.lower() not in VIDEO_SUFFIXES]

    if args.shared:
        for original in resources:
            text_file = convert(original, args.dest_dir)
            print(f"shared: {original.name} -> {text_file.relative_to(PROJECT_ROOT)}")
        return 0

    shared = shared_index()
    rows = []
    for original in resources:
        digest = md5(original)
        if digest in shared:
            shared_copy = shared[digest]
            text_name = f"{shared_copy.name}.md" if shared_copy.suffix.lower() in (".pdf", ".xlsx") else shared_copy.name
            rows.append((original.name, f"compartido: `inputs/curso/recursos/{shared_copy.name}` (idéntico, md5 {digest[:8]})",
                         f"`inputs/curso/recursos/{text_name}`"))
            print(f"shared already: {original.name}")
            continue
        text_file = convert(original, args.dest_dir / "recursos")
        rows.append((original.name, f"`recursos/{original.name}`", f"`recursos/{text_file.name}`"))
        print(f"converted: {original.name} -> {text_file.relative_to(PROJECT_ROOT)}")

    lines = [f"# Fuente — Sesión {args.session}", "", f"- Carpeta de origen: `{args.source_dir}`",
             f"- Convertido: {dt.date.today().isoformat()}"]
    for video in videos:
        lines.append(f"- Video: `{video}` · duración {hms(probe_duration(video))}")
    lines += ["", "## Recursos", "", "| Archivo | Original | Texto |", "|---|---|---|"]
    lines += [f"| {name} | {where} | {text} |" for name, where, text in rows]
    args.dest_dir.mkdir(parents=True, exist_ok=True)
    (args.dest_dir / "fuente.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"wrote {args.dest_dir / 'fuente.md'}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
